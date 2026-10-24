from pathlib import Path
from datetime import datetime
import re
import json
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

OUTPUT = Path(__file__).resolve().parent / "excels"
OUTPUT.mkdir(parents=True, exist_ok=True)

RUTA_IMAGEN = Path(__file__).resolve().parent / "Imagen" / "marca.png"

_ultimo_archivo = None
_ruta_archivo_real = None


def obtener_prompt(tema: str) -> str:
    if "editando un archivo" in tema.lower() or "modificación" in tema.lower():
        return (
            "Sos un experto en Excel. El usuario está editando un archivo que ya generaste en esta conversación.\n\n"
            "Revisá el historial, encontrá el JSON del Excel que generaste anteriormente "
            "y aplicá SOLO los cambios que pide el usuario.\n\n"
            "Devolvé EXCLUSIVAMENTE un objeto JSON válido, sin markdown, con esta estructura:\n"
            '{"nombre": "nombre_archivo", "hojas": [{"nombre": "Hoja1", '
            '"encabezados": ["Col1", "Col2"], "filas": [["a", 1], ["b", 2]]}]}\n\n'
            "REGLAS:\n"
            "- MANTENÉ los datos originales y SOLO agregá, quitá o modificá lo que el usuario pide.\n"
            "- Si el usuario pide agregar una columna, agregala a los encabezados y a todas las filas.\n"
            "- Devolvé el JSON COMPLETO y actualizado.\n"
            "- NO uses ```json. Solo el JSON puro.\n\n"
            f"PEDIDO DEL USUARIO:\n{tema}"
        )
    
    return (
        "Sos un experto en Excel. El usuario pidió un archivo EXCEL.\n\n"
        "Devolvé EXCLUSIVAMENTE un objeto JSON válido, sin markdown, sin explicaciones. "
        "Solo devolvé el JSON crudo con esta estructura exacta:\n"
        '{"nombre": "nombre_archivo", "hojas": [{"nombre": "Hoja1", '
        '"encabezados": ["Col1", "Col2"], "filas": [["a", 1], ["b", 2]]}]}\n\n'
        "REGLAS:\n"
        "- 'hojas' es una lista de 1 a 5 hojas.\n"
        "- 'encabezados' es la fila 1 de cada hoja.\n"
        "- 'filas' son los datos (desde fila 2). Usá null para celdas vacías.\n"
        "- Si el usuario no da datos exactos, INVENTÁ datos realistas coherentes con el pedido.\n"
        "- NO devuelvas código Python. NO uses ```json. Solo el JSON puro.\n\n"
        f"PEDIDO DEL USUARIO:\n{tema}"
    )


def construir_codigo(datos: dict) -> str:
    nombre = str(datos.get('nombre') or 'reporte')
    hojas = datos.get('hojas') or []

    if not isinstance(hojas, list) or not hojas:
        raise ValueError("La IA no devolvió el formato esperado.")

    # Limpiamos strings de cualquier caracter corrupto
    nombre = nombre.encode('utf-8', 'ignore').decode('utf-8', 'ignore')
    
    for hoja in hojas:
        if isinstance(hoja, dict):
            for key in ('nombre',):
                if hoja.get(key):
                    hoja[key] = str(hoja[key]).encode('utf-8', 'ignore').decode('utf-8', 'ignore')
            if hoja.get("encabezados"):
                hoja["encabezados"] = [str(e).encode('utf-8', 'ignore').decode('utf-8', 'ignore') for e in hoja["encabezados"]]
            if hoja.get("filas"):
                for fila in hoja["filas"]:
                    for i, v in enumerate(fila):
                        if isinstance(v, str):
                            fila[i] = v.encode('utf-8', 'ignore').decode('utf-8', 'ignore')

    json_str = json.dumps(hojas, ensure_ascii=False).replace("null", "None")

    return (
        f"nombre = {json.dumps(nombre)}\n"
        f"hojas = {json_str}\n"
        "crear(nombre=nombre, hojas=hojas)\n"
    )


def _slug(texto, maximo=60):
    s = re.sub(r"[^a-zA-Z0-9_\-]+", "_", str(texto or "reporte")).strip("_")
    return (s or "reporte")[:maximo]


def crear(nombre="reporte", hojas=None, **kwargs):
    global _ultimo_archivo, _ruta_archivo_real

    if not hojas or not isinstance(hojas, list):
        raise ValueError("Se requiere 'hojas' como lista no vacía")

    logo_b64 = ""
    if RUTA_IMAGEN.exists():
        with open(RUTA_IMAGEN, "rb") as f:
            logo_b64 = base64.b64encode(f.read()).decode("utf-8")

    nombre_archivo = _slug(nombre)
    fecha = datetime.now().strftime("%Y%m%d_%H%M%S")

    # ============ 1. GENERAR EXCEL ============
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for i, spec in enumerate(hojas):
        if not isinstance(spec, dict):
            continue
        ws = wb.create_sheet(title=str(spec.get("nombre") or f"Hoja{i+1}")[:31])

        encabezados = spec.get("encabezados") or []
        filas = spec.get("filas") or []
        num_columnas = len(encabezados) if encabezados else 1

        # ---- 1.1 COLUMNAS DEL LOGO (A y B) ----
        if i == 0:
            ws.column_dimensions['A'].width = 12
            ws.column_dimensions['B'].width = 12
            for r in range(1, 11):
                ws.row_dimensions[r].height = 20

        # ---- 1.2 INSERTAR LOGO EN A1:B10 ----
        if RUTA_IMAGEN.exists() and i == 0:
            try:
                ws.merge_cells("A1:B10")
                img = XLImage(str(RUTA_IMAGEN))
                img.width = 150
                img.height = 200
                ws.add_image(img, 'A1')
            except Exception as e:
                print(f"No se pudo cargar el logo en Excel: {e}")

        # ---- 1.3 CONFIGURAR COLUMNAS DE LA TABLA (Empiezan en C) ----
        offset_col = 3
        for col_idx in range(num_columnas):
            col_letter = get_column_letter(col_idx + offset_col)
            if encabezados and col_idx < len(encabezados):
                ws.column_dimensions[col_letter].width = max(15, len(str(encabezados[col_idx])) + 8)
            else:
                ws.column_dimensions[col_letter].width = 15

        # ---- 1.4 ESCRIBIR ENCABEZADOS ----
        fila_actual_tabla = 1
        for col_idx, enc in enumerate(encabezados):
            col_actual = col_idx + offset_col
            cell = ws.cell(row=fila_actual_tabla, column=col_actual, value=enc)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", start_color="17365D")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # ---- 1.5 ESCRIBIR FILAS ----
        for r_idx, fila in enumerate(filas, start=fila_actual_tabla + 1):
            for c_idx, val in enumerate(fila):
                cell = ws.cell(row=r_idx, column=c_idx + offset_col, value=val)
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ruta_xlsx = OUTPUT / f"{nombre_archivo}_{fecha}.xlsx"
    wb.save(ruta_xlsx)

    # ============ 2. GENERAR HTML (Layout idéntico) ============
    html_content = """
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body { font-family: 'Calibri', 'Segoe UI', sans-serif; font-size: 11pt; margin: 20px; background: #fff; color: #333; }
            .dashboard-layout { display: flex; align-items: flex-start; gap: 20px; }
            .logo-sidebar { flex: 0 0 150px; display: flex; align-items: center; justify-content: center; height: 200px; }
            .logo-img { max-width: 100%; max-height: 100%; object-fit: contain; }
            .table-content { flex: 1; overflow-x: auto; }
            table { border-collapse: collapse; width: 100%; }
            th { background-color: #17365D; color: #FFFFFF; font-weight: bold; border: 1px solid #a0a0a0; padding: 8px; text-align: center; }
            td { border: 1px solid #d0d0d0; padding: 6px 8px; text-align: center; }
            tr:nth-child(even) td { background-color: #f9f9f9; }
        </style>
    </head>
    <body>
    """
    
    for i, spec in enumerate(hojas):
        html_content += '<div class="dashboard-layout">'
        
        if i == 0 and logo_b64:
            html_content += f'<div class="logo-sidebar"><img src="data:image/png;base64,{logo_b64}" class="logo-img"></div>'
        else:
            html_content += '<div class="logo-sidebar"></div>'

        html_content += '<div class="table-content">'
        
        if len(hojas) > 1:
            html_content += f"<h3 style='margin-top:0; color:#17365D;'>{spec.get('nombre', f'Hoja {i+1}')}</h3>"

        encabezados = spec.get("encabezados") or []
        filas = spec.get("filas") or []

        html_content += "<table>\n"
        if encabezados:
            html_content += "<tr>" + "".join(f"<th>{e}</th>" for e in encabezados) + "</tr>\n"
        for fila in filas:
            html_content += "<tr>" + "".join(f"<td>{v if v is not None else ''}</td>" for v in fila) + "</tr>\n"
        html_content += "</table>\n"
        
        html_content += '</div></div><br><br>'

    html_content += "</body></html>"

    ruta_html = OUTPUT / f"{nombre_archivo}_{fecha}.html"
    ruta_html.write_text(html_content, encoding="utf-8", errors="ignore")

    _ultimo_archivo = str(ruta_html)
    _ruta_archivo_real = str(ruta_xlsx)

    print(f"[excel] HTML generado: {ruta_html}")
    print(f"[excel] Excel generado: {ruta_xlsx}")
    return str(ruta_html)


def ultimo_archivo():
    return _ultimo_archivo


def archivo_real():
    return _ruta_archivo_real